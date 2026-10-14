"""Все карточки MITEX + contact-scraper, контроль покрытия и выгрузки."""
import argparse
import copy
import csv
import json
import logging
import re
import sys
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src.services.catalog_client import collect_catalog, collect_detail, normalize_website, website_urls
from src.services.http_client import ROOT
from src.vendor import contact_scraper as skill

logging.getLogger('httpx').setLevel(logging.WARNING)
logging.getLogger('contact-scraper').setLevel(logging.WARNING)
SITE=ROOT/'data/site'
STATE=ROOT/'data/contacts'
STATE.mkdir(parents=True,exist_ok=True)
local=threading.local()
NETWORKS={'telegram':'Telegram','max':'MAX','whatsapp':'WhatsApp','vk':'VK','instagram':'Instagram','facebook':'Facebook','ok':'OK','linkedin':'LinkedIn','youtube':'YouTube','x':'X','wechat':'WeChat','line':'Line','xing':'Xing'}
original_phone=skill.fmt_phone
def phone(raw):
    digits=re.sub(r'\D','',raw)
    if len(digits)==11 and digits.startswith('8'): raw='+7'+digits[1:]
    elif len(digits)==11 and digits.startswith('7'): raw='+'+digits
    return original_phone(raw)
skill.fmt_phone=phone

def now(): return datetime.now(timezone.utc).isoformat()
def write(path,value):
    temp=path.with_suffix('.tmp')
    temp.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    temp.replace(path)
def domain(url): return (urlparse(skill.normalize_url(url)).hostname or '').lower().removeprefix('www.')

# Строгие домены вместо поиска подстроки "t.me" в произвольной ссылке.
skill.SOCIAL_MATCHERS['telegram']=lambda h,host: host.removeprefix('www.') in ('t.me','telegram.me','telegram.dog') and not urlparse(h).path.startswith(('/share','/iv','/addstickers','/addemoji','/proxy','/socks'))
original_extract=skill.extract_contacts
def extract(page_url,soup,acc):
    for a in soup.find_all('a',href=True):
        if not a['href'].startswith(('tel:','mailto:')): a['href']=urljoin(page_url,a['href'])
    page_acc={key:set() for key in skill.CONTACT_KEYS}
    original_extract(page_url,soup,page_acc)
    for kind,values in page_acc.items():
        for value in values:
            if not value: continue
            acc[kind].add(value)
            local.sources.append({'kind':kind,'value':value,'source_url':page_url})
skill.extract_contacts=extract

original_fetch=skill.fetch
def fetch(url,client,timeout,use_playwright):
    content=original_fetch(url,client,timeout,use_playwright)
    local.pages.append({'url':url,'success':bool(content)})
    return content
skill.fetch=fetch

def details(records,workers):
    todo=[x for x in records if x.get('detail_version')!=2 or x.get('detail_status')=='error']
    by_id={x['id']:i for i,x in enumerate(records)}
    def work(x):
        result=copy.deepcopy(x)
        try: collect_detail(result)
        except Exception as error:
            result.update(detail_status='error',detail_note=str(error),detail_checked_at=now())
        return result
    print(f'DETAILS {len(todo)} pending of {len(records)}',flush=True)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures=[pool.submit(work,x) for x in todo]
        for count,future in enumerate(as_completed(futures),1):
            x=future.result(); records[by_id[x['id']]]=x
            write(SITE/'exhibitors.json',records)
            if count%20==0 or count==len(todo):
                print(f'DETAILS {count}/{len(todo)}; websites={sum(bool(x.get("website")) for x in records)}',flush=True)
    return records

def sites(records,workers,timeout,max_pages,retry_unreachable=False,direct=False,retry_failed_pages=False):
    domains={}
    for x in records:
        for website in x.get('websites',[]) or ([x['website']] if x['website'] else []):
            key=domain(website)
            if key and key not in domains: domains[key]={'name':x['name'],'site':website}
    path=STATE/'contacts.jsonl'
    done={}
    if path.exists():
        for line in path.read_text(encoding='utf-8').splitlines():
            row=json.loads(line); done[domain(row['site'])]=row
    todo=[x for key,x in domains.items() if key not in done or retry_unreachable and not done[key]['reachable'] or retry_failed_pages and any(not p['success'] for p in done[key]['pages'])]
    skill.INTL_MODE=True
    def work(x):
        local.sources=[]; local.pages=[]
        try:
            row=skill.crawl_site(x['name'],x['site'],timeout,max_pages,False)
            row.pop('home_text',None)
            row['description']=row.get('meta_desc','')
        except Exception as error:
            row={**x,'home':skill.base_url(x['site']),'reachable':False,'error':str(error),**{key:[] for key in skill.CONTACT_KEYS}}
        row.update(checked_at=now(),contact_sources=local.sources,pages=local.pages,route='direct' if direct else 'system')
        return row
    print(f'SITES {len(domains)} domains; {len(todo)} pending',flush=True)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures=[pool.submit(work,x) for x in todo]
        for count,future in enumerate(as_completed(futures),1):
            row=future.result(); key=domain(row['site']); previous=done.get(key)
            if previous:
                attempt={'route':row['route'],'checked_at':row['checked_at'],'reachable':row['reachable'],'pages':row['pages']}
                row['attempts']=previous.get('attempts',[{'route':previous.get('route','system'),'checked_at':previous['checked_at'],'reachable':previous['reachable'],'pages':previous['pages']}])+[attempt]
                for field in skill.CONTACT_KEYS: row[field]=sorted(set(row[field]+previous[field]))
                row['contact_sources']=list({(e['kind'],e['value'],e['source_url']):e for e in previous['contact_sources']+row['contact_sources']}.values())
                row['reachable']=row['reachable'] or previous['reachable']
                pages={p['url']:p for p in previous['pages']}
                for p in row['pages']:
                    if p['success'] or p['url'] not in pages: pages[p['url']]=p
                row['pages']=list(pages.values())
            done[key]=row
            with path.open('a',encoding='utf-8') as file: file.write(json.dumps(row,ensure_ascii=False)+'\n')
            if count%10==0 or count==len(todo):
                print(f'SITES {count}/{len(todo)}; reachable={sum(x["reachable"] for x in done.values())}; Telegram domains={sum(bool(x.get("telegram")) for x in done.values())}',flush=True)
    return {key:row for key,row in done.items() if key in domains}

def finalize(records,results):
    active={domain(url) for x in records for url in x.get('websites',[]) if domain(url)}
    results={key:row for key,row in results.items() if key in active}
    for x in records:
        rows=[results[domain(url)] for url in (x.get('websites',[]) or ([x['website']] if x['website'] else [])) if domain(url) in results]
        if not x['website']:
            x['enrichment_status']='no_website'
        elif not rows:
            x['enrichment_status']='not_collected'
        else:
            x['enrichment_status']='collected' if any(r['reachable'] for r in rows) else 'unreachable'
        x['website_checks']=[{'website':r['site'],'reachable':r['reachable'],'pages':r['pages'],'checked_at':r['checked_at'],'route':r.get('route','system')} for r in rows]
        sources=x.setdefault('contact_sources',[])
        social={(s['network'],s['url']):s for s in x['contacts']['social']}
        for row in rows:
            for kind in ('emails','phones'): x['contacts'][kind]=sorted(set(x['contacts'][kind]+row[kind]))
            for evidence in row['contact_sources']:
                if evidence not in sources: sources.append(evidence)
            for kind,network in NETWORKS.items():
                for url in row[kind]:
                    source=next((e['source_url'] for e in row['contact_sources'] if e['kind']==kind and e['value']==url),row['site'])
                    social[(network,url)]={**social.get((network,url),{}),'network':network,'url':url,'source_url':source}
        x['contacts']['social']=list(social.values())
    write(SITE/'exhibitors.json',records)
    meta=json.loads((SITE/'meta.json').read_text(encoding='utf-8'))
    meta.update(checked_at=now(),catalog_count=len(records),
        details_collected=sum(x['detail_status']=='collected' for x in records),
        details_checked=sum(bool(x.get('detail_checked_at')) for x in records),
        sites_enriched=sum(x['enrichment_status']=='collected' for x in records),
        website_domains=len(results),sites_attempted=len(results),sites_reachable=sum(x['reachable'] for x in results.values()),
        companies_with_website=sum(bool(x['website']) for x in records),
        companies_without_website=sum(not x['website'] for x in records),
        details_status=dict(Counter(x['detail_status'] for x in records)),
        enrichment_status=dict(Counter(x['enrichment_status'] for x in records)),
        errors=[{'id':x['id'],'name':x['name'],'error':x.get('detail_note','')} for x in records if x['detail_status']=='error'])
    telegram={s['url'].rstrip('/') for x in records for s in x['contacts']['social'] if s['network']=='Telegram'}
    meta['telegram_unique_links']=len(telegram)
    meta['telegram_companies']=sum(any(s['network']=='Telegram' for s in x['contacts']['social']) for x in records)
    meta['contact_collector']='contact-scraper skill; full catalog; main + up to 6 contact/about pages per domain'
    meta['site_pages_requested']=sum(len(row['pages']) for row in results.values())
    meta['site_pages_successful']=sum(page['success'] for row in results.values() for page in row['pages'])
    meta['sites_partial']=sum(row['reachable'] and any(not p['success'] for p in row['pages']) for row in results.values())
    meta['direct_attempts']=sum(row.get('route')=='direct' for row in results.values())
    meta['recovered_without_proxy']=sum(bool(row.get('route')=='direct' and row['reachable'] and row.get('attempts') and not row['attempts'][0]['reachable']) for row in results.values())
    write(SITE/'meta.json',meta)
    report={'checked_at':now(),'unreachable_sites':[{'website':row['site'],'company':row['name'],'pages':row['pages'],'error':row.get('error','Главная страница не получена')} for row in results.values() if not row['reachable']],
        'failed_pages':[{'website':row['site'],'url':page['url']} for row in results.values() for page in row['pages'] if not page['success']],
        'without_website':[{'id':x['id'],'name':x['name'],'source_url':x['source_url']} for x in records if not x['website']]}
    write(SITE/'contact_report.json',report)
    headers=['Компания','Страна','Стенд','Сайт','Email','Телефон','Telegram','MAX','WhatsApp','Соцсети','Источник','Статус карточки','Статус сайта','Рубрики','Описание']
    table=[]
    for x in records:
        c=x['contacts']
        network=lambda name:', '.join(s['url'] for s in c['social'] if s['network']==name)
        table.append([x['name'],x['country'],', '.join(x['stands']),', '.join(x.get('websites',[]) or ([x['website']] if x['website'] else [])),', '.join(c['emails']),', '.join(c['phones']),network('Telegram'),network('MAX'),network('WhatsApp'),', '.join(s['url'] for s in c['social']),x['source_url'],x['detail_status'],x['enrichment_status'],', '.join(x.get('categories',[])),x['about']])
    with (SITE/'exhibitors.csv').open('w',encoding='utf-8-sig',newline='') as file:
        writer=csv.writer(file,delimiter=';');writer.writerow(headers);writer.writerows(table)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='Участники';ws.append(headers)
    for row in table: ws.append(["'"+v if isinstance(v,str) and v.startswith('=') else v for v in row])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='182A2E')
    for index in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(index)].width=55 if index in (10,14,15) else 28
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(SITE/'exhibitors.xlsx')
    print('FINAL',json.dumps({k:meta[k] for k in ('catalog_count','details_checked','details_status','companies_with_website','website_domains','sites_reachable','telegram_unique_links','telegram_companies')},ensure_ascii=False),flush=True)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--stage',choices=['all','details','sites','finalize'],default='all');parser.add_argument('--workers',type=int,default=6);parser.add_argument('--timeout',type=int,default=15);parser.add_argument('--max-pages',type=int,default=6)
    parser.add_argument('--direct',action='store_true',help='Обойти системный прокси только в HTTP-клиенте сайтов')
    parser.add_argument('--retry-unreachable',action='store_true',help='Повторить недоступные сайты')
    parser.add_argument('--retry-failed-pages',action='store_true',help='Повторить сайты с неудачными контактными страницами')
    args=parser.parse_args()
    if args.direct:
        original_client=skill.httpx.Client
        def direct_client(*a,**kw):
            kw['trust_env']=False
            return original_client(*a,**kw)
        skill.httpx.Client=direct_client
    records=json.loads((SITE/'exhibitors.json').read_text(encoding='utf-8'))
    for x in records:
        original=x.get('websites',[]) or ([x['website']] if x['website'] else [])
        fixed=[value for url in original for value in website_urls(url)]
        if fixed!=original: x['website_corrections']={'original':original,'normalized':fixed}
        x['websites']=fixed
        x['website']=fixed[0] if fixed else ''
    if args.stage in ('all','details'):
        fresh=collect_catalog(refresh=True)
        old={x['id']:x for x in records};records=[old.get(x['id'],x) for x in fresh]
        records=details(records,args.workers)
    if args.stage in ('all','sites'): results=sites(records,args.workers,args.timeout,args.max_pages,args.retry_unreachable,args.direct,args.retry_failed_pages)
    else:
        results={}
        if (STATE/'contacts.jsonl').exists():
            for line in (STATE/'contacts.jsonl').read_text(encoding='utf-8').splitlines():
                row=json.loads(line);results[domain(row['site'])]=row
    finalize(records,results)

if __name__=='__main__': main()
