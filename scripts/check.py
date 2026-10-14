"""Содержательные проверки целостности снимка и автономной сборки."""
import csv
import json
from pathlib import Path
from bs4 import BeautifulSoup

ROOT=Path(__file__).resolve().parents[1]
records=json.loads((ROOT/'data/site/exhibitors.json').read_text(encoding='utf-8'))
meta=json.loads((ROOT/'data/site/meta.json').read_text(encoding='utf-8'))
assert records and len(records)==meta['catalog_count']
assert len({r['id'] for r in records})==len(records), 'Дубли id'
assert len({r['source_url'] for r in records})==len(records), 'Дубли URL'
assert all(r['name'].strip() and '\ufffd' not in r['name']+r['about'] for r in records), 'Повреждённая кодировка'
assert sum(r['detail_status']=='collected' for r in records)==meta['details_collected']
with (ROOT/'site/exhibitors.csv').open(encoding='utf-8-sig',newline='') as file:
    rows=list(csv.reader(file,delimiter=';'))
assert len(rows)-1==len(records)
html=BeautifulSoup((ROOT/'site/index.html').read_text(encoding='utf-8'),'html.parser')
for tag in html.find_all(['script','link']):
    target=tag.get('src') or tag.get('href')
    if target and not target.startswith('http'): assert (ROOT/'site'/target.split('?')[0]).is_file(),target
js=(ROOT/'site/data.js').read_text(encoding='utf-8')
payload=json.loads(js.removeprefix('window.MITEX_DATA = ').rstrip(';\n'))
assert len(payload['exhibitors'])==len(records)
events=payload['program']['events']
assert events and payload['program']['status']=='collected'
assert len({e['id'] for e in events})==len(events)
assert all(e['date'].startswith('2026-11-') for e in events), 'Чужой год программы'
assert all('/2026/' in e['source_url'] or e['id']=='press-14229' for e in events)
assert len({p['id'] for p in payload['people']})==len(payload['people'])
import re
for person in payload['people']:
    for social in person['social']:
        if social.get('affiliation_source'):
            employer=social['employer_match']
            assert re.search(r'(?<![а-яёa-z])'+re.escape(employer)+r'(?![а-яёa-z])',person['role'].lower()), 'Неверная привязка работодателя'
from collections import Counter
counts=Counter(p['channel_url'] for p in payload['feed'])
assert counts and all(value==5 for value in counts.values()), 'Нужно пять постов на канал'
assert len(counts)==payload['feed_report']['channels']
assert all(len(p['text'].split())<=24 for p in payload['feed']), 'В ленте только краткие выдержки; полный текст — в оригинальном плеере'
assert len({p['url'] for p in payload['feed']})==len(payload['feed'])
assert all(v['date'].startswith('2026') for v in payload['videos'] if v['kind']=='invitation')
assert all(v['date'].startswith('2025') for v in payload['videos'] if v['kind']=='archive')
assert len(payload['novelty_registry'])==len(records)
assert not any(i['shown_at_exhibition'] for i in payload['items']), 'Показ будущей выставки нельзя подтверждать заранее'
assert not any('???' in (ROOT/'site'/f).read_text(encoding='utf-8') for f in ('portal.js','index.html','novinki.html'))
if meta.get('details_checked'):
    assert meta['details_checked']==len(records), 'Не завершён обход всех карточек'
    assert all(r['detail_status'] in ('collected','no_contacts','error') for r in records)
    assert sum(bool(r['website']) for r in records)==meta['companies_with_website']
    assert all(r['enrichment_status'] in ('collected','unreachable','no_website') for r in records), 'Не завершён обход сайтов'
    assert sum(any(s['network']=='Telegram' for s in r['contacts']['social']) for r in records)==meta['telegram_companies']
    assert all(s.get('source_url') for r in records for s in r['contacts']['social']), 'Соцсеть без источника'
    from openpyxl import load_workbook
    workbook=load_workbook(ROOT/'site/exhibitors.xlsx',read_only=True,data_only=False)
    assert workbook.active.max_row==len(records)+1
    assert not any(cell.data_type=='f' for row in workbook.active for cell in row), 'Нежелательные формулы в контактах'
    workbook.close()
print(f'PASS: {len(records)} unique records; UTF-8; coverage; CSV; assets; embedded data; program status')
